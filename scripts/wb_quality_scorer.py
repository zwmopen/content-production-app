# -*- coding: utf-8 -*-
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(errors='replace')
"""
成品质量打分器（网页CDP产线 · WorkBuddy）
==========================================
只读扫描成品库，逐套打分，输出 xlsx 打分表。绝不改动/移动/删除任何成品文件。
垃圾候选只标记不处置（用户铁律：垃圾就是垃圾，不抢救）。

【检测标准 / 评分表 · 总分 100】
A 图片 40 分
  A1 数量 20：成品图 ≥4 张=20；3 张=12；2 张=6；1 张=0（主脑铁律：≥2 原料必须出 ≥4 图）
  A2 竖屏 10：Pillow 可打开且 h/w ≥1.25 的比例 = 通过率×10
  A3 清晰度 10：平均宽度 ≥1000px=6，≥800=3；平均单图 ≥300KB=4，≥150KB=2
B 文案 40 分（以 三平台文案.txt 为准，缺失则用 文案.txt/小红书文案.txt 降档）
  B1 存在 10：三平台文案.txt=10；仅其他文案=6；全无=0
  B2 槽位 10：<<<XHS_START>>> / <<<XHS_2_START>>> / <<<DOUYIN_START>>> 各 3~4 分，全齐 10
  B3 实质字数 10：剔除 <<<...>>>、空白、盲文空格后 ≥300 字=10；200-299=6；<200=0
  B4 卫生 10：无占位行（[主标题]【种草正文】等写作指令）=3；无「标题：/正文：/话题：」标签外泄=3；
     盲文占位行（\\u2800）≥3 行=2；无相邻重复行=2
C 完整性 20 分
  C1 manifest 10：存在且 status=PASS=10；存在非 PASS=5；无=0
  C2 图数一致 5：manifest.imageCount == 实际图数
  C3 附属文案 5：小红书文案.txt / 文案.txt 任一存在

【分级】≥85 优 · 70-84 良 · 60-69 合格 · <60 ⚠️垃圾候选
"""
import os, re, json, sys
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

POOL = r'D:\AICode\项目推进\projects\江湖有旅人\主项目\成品库（GPT+本地脚本制作）'
OUT_XLSX = r'D:\AICode\运行数据\江湖有旅人\成品质量打分表.xlsx'
IMG_EXT = {'.png', '.jpg', '.jpeg', '.webp'}
PLACEHOLDER_LINE = re.compile(r'^[\[【][^\]】]{2,}[\]】]$')
PLACEHOLDER_HINT = ('主标题', '种草正文', '话题标签', '方案决策版', '口播脚本', '大纲', '占位', '待填',
                    '请填', '留资号召', '预算参考', '决策亮点', '服务保障', '空话', '亮点提炼',
                    '详细行程', '写作', '文案模板', '此处')
LABEL_LEAK = re.compile(r'^(标题|正文|话题)[:：]')
SKIP_DIRS = {'_垃圾作品样本', '_不合格成品合集', '_异常素材（脚本失败隔离）', '_重复待处理',
             '原素材', '原图', '原料', '素材', '素材图', '配图'}


def braille_count(text):
    return sum(1 for ln in text.splitlines() if ln.strip() == '\u2800')


def substance(text):
    t = re.sub(r'<<<[^>]*>>>', '', text)
    return len(re.sub(r'[\s\u2800]+', '', t))


def find_package_dirs(root):
    hits = []
    for dp, dn, fn in os.walk(root):
        dn[:] = [d for d in sorted(dn) if d not in SKIP_DIRS]
        imgs = [f for f in fn if os.path.splitext(f)[1].lower() in IMG_EXT]
        has_copy = any(f in ('三平台文案.txt', '文案.txt', '小红书文案.txt', '小红书发布文案.txt') for f in fn)
        if imgs and (has_copy or 'manifest.json' in fn):
            hits.append(dp)
    return hits


def score_package(d):
    r = {'dir': d, 'name': os.path.basename(d), 'imgs': 0, 'vpass': '', 'wavg': 0, 'kbavg': 0,
         'A': 0, 'B': 0, 'C': 0, 'notes': []}
    files = os.listdir(d)
    imgs = [os.path.join(d, f) for f in files if os.path.splitext(f)[1].lower() in IMG_EXT]
    r['imgs'] = len(imgs)
    # A1 数量
    n = len(imgs)
    r['A'] += 20 if n >= 4 else 12 if n == 3 else 6 if n == 2 else 0
    if n < 4:
        r['notes'].append(f'图仅{n}张(<4)')
    # A2/A3 Pillow
    ok = 0; ws = []; kbs = []
    for p in imgs:
        try:
            with Image.open(p) as im:
                w, h = im.size
                if h / max(w, 1) >= 1.25: ok += 1
                ws.append(w)
            kbs.append(os.path.getsize(p) / 1024)
        except Exception as e:
            r['notes'].append('坏图:' + os.path.basename(p)[:20])
    if n:
        rate = ok / n
        r['vpass'] = f'{ok}/{n}'
        r['A'] += round(rate * 10)
        if rate < 1: r['notes'].append(f'竖屏通过率{ok}/{n}')
        wavg = sum(ws) / len(ws) if ws else 0
        kbavg = sum(kbs) / len(kbs) if kbs else 0
        r['wavg'] = int(wavg); r['kbavg'] = int(kbavg)
        r['A'] += 6 if wavg >= 1000 else 3 if wavg >= 800 else 0
        r['A'] += 4 if kbavg >= 300 else 2 if kbavg >= 150 else 0
    # B 文案
    main = os.path.join(d, '三平台文案.txt')
    if os.path.exists(main):
        r['B'] += 10; src = main
    elif os.path.exists(os.path.join(d, '文案.txt')):
        r['B'] += 6; src = os.path.join(d, '文案.txt'); r['notes'].append('缺三平台文案')
    elif os.path.exists(os.path.join(d, '小红书文案.txt')):
        r['B'] += 6; src = os.path.join(d, '小红书文案.txt'); r['notes'].append('缺三平台文案')
    else:
        src = None; r['notes'].append('无任何文案')
    if src:
        try: text = open(src, 'r', encoding='utf-8-sig', errors='replace').read()
        except Exception: text = ''
        s = 0
        for mark in ('<<<XHS_START>>>', '<<<XHS_2_START>>>', '<<<DOUYIN_START>>>'):
            if mark in text: s += 3.34
        s = min(10, round(s))
        r['B'] += s
        if s < 10: r['notes'].append(f'槽位不全({s}/10)')
        sub = substance(text)
        r['B'] += 10 if sub >= 300 else 6 if sub >= 200 else 0
        if sub < 300: r['notes'].append(f'实质字数{sub}')
        h = 0
        lines = text.splitlines()
        ph = sum(1 for ln in lines if PLACEHOLDER_LINE.match(ln.strip()) or any(k in ln for k in PLACEHOLDER_HINT))
        leak = sum(1 for ln in lines if LABEL_LEAK.match(ln.strip()))
        bc = braille_count(text)
        dup = sum(1 for a, b in zip(lines, lines[1:]) if a.strip() and a.strip() == b.strip())
        if ph == 0: h += 3
        else: r['notes'].append(f'占位行{ph}')
        if leak == 0: h += 3
        else: r['notes'].append(f'标签外泄{leak}行')
        if bc >= 3: h += 2
        else: r['notes'].append(f'盲文占位仅{bc}行')
        if dup == 0: h += 2
        else: r['notes'].append(f'相邻重复{dup}行')
        r['B'] += h
    # C 完整性
    mf = os.path.join(d, 'manifest.json')
    if os.path.exists(mf):
        try:
            m = json.load(open(mf, encoding='utf-8'))
            if str(m.get('status', '')).upper() == 'PASS': r['C'] += 10
            else: r['C'] += 5; r['notes'].append('manifest非PASS')
            if int(m.get('imageCount', -1)) == n: r['C'] += 5
            else: r['notes'].append(f"imageCount{m.get('imageCount')}≠实际{n}")
        except Exception:
            r['C'] += 0; r['notes'].append('manifest损坏')
    else:
        r['notes'].append('无manifest')
    if os.path.exists(os.path.join(d, '小红书文案.txt')) or os.path.exists(os.path.join(d, '文案.txt')):
        r['C'] += 5
    else:
        r['notes'].append('缺附属文案')
    r['total'] = r['A'] + r['B'] + r['C']
    r['grade'] = '优' if r['total'] >= 85 else '良' if r['total'] >= 70 else '合格' if r['total'] >= 60 else '⚠️垃圾候选'
    return r


def main():
    dirs = find_package_dirs(POOL)
    rows = [score_package(d) for d in dirs]
    rows.sort(key=lambda x: x['total'])
    wb = Workbook(); ws = wb.active; ws.title = '打分表'
    header = ['总分', '分级', '作品名', '图数', '竖屏通过', '均宽px', '均KB', '图分/40', '文案分/40', '完整性/20', '问题明细', '路径']
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='4472C4')
        c.alignment = Alignment(horizontal='center')
    red = PatternFill('solid', fgColor='FFC7CE'); yellow = PatternFill('solid', fgColor='FFEB9C')
    for r in rows:
        ws.append([r['total'], r['grade'], r['name'], r['imgs'], r['vpass'], r['wavg'], r['kbavg'],
                   r['A'], r['B'], r['C'], '；'.join(r['notes']), r['dir']])
        if r['total'] < 60:
            for c in ws[ws.max_row]: c.fill = red
        elif r['total'] < 70:
            for c in ws[ws.max_row]: c.fill = yellow
    ws.column_dimensions['C'].width = 52; ws.column_dimensions['K'].width = 44; ws.column_dimensions['L'].width = 60
    ws.freeze_panes = 'A2'
    wb.save(OUT_XLSX)
    bad = [r for r in rows if r['total'] < 60]
    print('扫描成品套数:', len(rows))
    from collections import Counter
    print('分级分布:', dict(Counter(r['grade'] for r in rows)))
    print('垃圾候选:', len(bad))
    for r in bad[:20]:
        print(f"  [{r['total']:3d}] {r['name'][:56]} | {'；'.join(r['notes'])[:90]}")
    print('打分表 ->', OUT_XLSX)


if __name__ == '__main__':
    main()
